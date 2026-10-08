"""Update the reproduction section from saved results and add a conclusion."""
from pathlib import Path
import json
import statistics
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'ReAct_ Synergizing Reasoning and Acting (HotpotQA Reproduction) - slides3-5 revised.pptx'
OUTPUT = ROOT / 'ReAct_HotpotQA_Reproduction_final.pptx'
prs = Presentation(SOURCE)
NAVY, BLUE, GREEN, GRAY = '0D1326', '235BC4', '147D69', '64748B'
notes = []

def text(s,x,y,w,h,value,size=18,color=NAVY,bold=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Noto Sans SC'; p.font.size=Pt(size)
        p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(6)
    return box

def panel(s,x,y,w,h,color='F0F4FA'):
    box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb=RGBColor.from_string(color)
    box.line.fill.background()
    return box

def base(index,kicker,title,subtitle):
    s=prs.slides[index] if index<len(prs.slides) else prs.slides.add_slide(prs.slide_layouts[6])
    for x in list(s.shapes):x._element.getparent().remove(x._element)
    s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('FFFFFF')
    text(s,1.11,.44,11.11,.33,f'{index+1:02d} · {kicker}',14,BLUE,True)
    text(s,1.11,.88,11.11,.65,title,27,NAVY,True)
    text(s,1.11,1.55,11.11,.5,subtitle,16,GRAY)
    text(s,12.05,7.13,.4,.23,f'{index+1:02d}',10,GRAY)
    return s

def table(s,x,y,w,headers,rows,widths,rowheight=.65,highlight=None,size=16):
    t=s.shapes.add_table(len(rows)+1,len(headers), Inches(x),Inches(y),Inches(w), Inches(rowheight*(len(rows)+1))).table
    for c,width in zip(t.columns,widths):c.width=Inches(width)
    for ri,vals in enumerate([headers]+rows):
        t.rows[ri].height=Inches(rowheight)
        for ci,val in enumerate(vals):
            cell=t.cell(ri,ci);cell.text=str(val)
            cell.margin_left=cell.margin_right=Inches(.13)
            cell.margin_top=cell.margin_bottom=Inches(.06)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(NAVY if ri==0 else 'E3F3ED' if ri-1==highlight else 'F0F4FA' if ri%2 else 'F8FAFC')
            p=cell.text_frame.paragraphs[0];p.font.name='Noto Sans SC';p.font.size=Pt(size)
            p.font.bold=ri==0 or ri-1==highlight
            p.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else NAVY)
    return t

def add_note(s,title,seconds,value):
    s.notes_slide.notes_text_frame.text=value
    notes.append(f'## {title}（约{seconds}秒）\n\n{value}\n')

def records(method,hybrid=False):
    path=ROOT/'results'/('qa-v6-hybrid-50' if hybrid else 'qa-v5-100')/'hotpotqa'/method/'episodes.jsonl'
    return [json.loads(line) for line in path.read_text(encoding='utf-8').splitlines() if line.strip()]

basic={m:records(m) for m in ['standard','cot','act','react']}
hybrid={m:records(m,True) for m in ['cot_sc','cot_sc_to_react','react_to_cot_sc']}
assert all(len(rs)==500 for rs in basic.values())
assert all(len(rs)==50 for rs in hybrid.values())
assert all([r['question_idx'] for r in rs]==[r['question_idx'] for r in basic['standard']] for rs in basic.values())
assert all([r['question_idx'] for r in rs]==[r['question_idx'] for r in basic['standard'][:50]] for rs in hybrid.values())

s=base(5,'Reproduction Design','Same Mechanism, Substitute Model','HotpotQA prompting reproduction • No fine-tuning • Question-only input')
table(s,1.11,2.22,11.11,['Controlled choice','Implementation'],[
    ['Model','GPT-4.1 mini · fixed snapshot 2025-04-14'],
    ['Prompts','Original six examples from the repository'],
    ['Samples','4 basic methods × 500 questions; 3 added methods × 50'],
    ['Fair comparison','Same shuffled question order · seed 233'],
    ['Generation','Temperature 0; seven interaction steps'],
    ['CoT-SC exception','21 samples per question · temperature 0.7'],
],[2.5,8.61],.61,size=17)
panel(s,1.11,6.64,11.11,.47,'E3F3ED')
text(s,1.3,6.73,10.73,.28,'Model replacement tests the method; it does not reproduce the original scores exactly.',16,GREEN,True)
add_note(s,'第6页：复现设计',45,'接下来介绍我们的复现。我们选择 HotpotQA，保留 question-only 设置和仓库原始的六个示例，不进行微调。模型换成固定版本的 GPT-4.1 mini，因此这是替代模型下的方法复现，不能要求分数与原论文完全一致。四种基本方法测试相同的五百题，新增三种方法测试其中的前五十题，随机种子固定为二百三十三。基本方法的 temperature 为零；CoT 自一致性按照论文设置，每题采样二十一条，temperature 为零点七。我们没有启用模型服务商的内置搜索，也没有配置额外的 reasoning 模式，但这不等于模型内部完全没有推理。')

s=base(6,'Engineering','The Script Executes; The Environment Observes','Standard / CoT: one-shot generation    |    Act / ReAct: an interaction loop')
steps=[('INPUT','Question + few-shot + history'),('MODEL','Generate Thought / Action'),('ENVIRONMENT','Execute Search / Lookup'),('FEEDBACK','Append Observation; repeat')]
for i,(label,content) in enumerate(steps):
    y=2.2+i*.91
    panel(s,1.11,y,4.32,.72)
    text(s,1.3,y+.1,3.94,.2,label,11,BLUE,True)
    text(s,1.3,y+.34,3.94,.29,content,16)
    if i<3:text(s,3.12,y+.72,.3,.2,'↓',13,GRAY)
table(s,5.87,2.2,6.35,['Script','Responsibility'],[
    ['run_qa.py','Prompts, scheduling, scoring summary'],
    ['run_hotpotqa.py','Model calls and action parsing'],
    ['wikienv.py\nwrappers.py','Wikipedia retrieval, tasks, metrics'],
    ['extend_qa.py','Extend sample target; resume runs'],
],[2.18,4.17],.72,size=15)
panel(s,1.11,6.19,11.11,.65,'E3F3ED')
text(s,1.3,6.31,10.73,.35,'Saved per question: raw output · full trajectory · EM / F1 · token usage',17,GREEN,True)
text(s,1.11,6.98,11.11,.25,'Adaptations: chat instructions; 256-token interactions / 512-token CoT; logged Unicode and memory fixes.',11,GRAY)
add_note(s,'第7页：工程实现',45,'我们以原始 notebook 为参考，把实验组织成可批量执行的脚本。run_qa 负责构造提示、安排样本和汇总结果；run_hotpotqa 负责模型调用和动作解析；wikienv 执行 Wikipedia 请求，wrappers 加载任务并计算评分。模型决定动作，环境返回观察，再把完整历史送入下一轮。Standard 和 CoT 不进入检索循环。我们先检查小样本，再跑一百题，通过 extend_qa 扩展到五百题并跳过已完成样本。每题保存输出、轨迹和用量。为适配聊天接口，我们调整了格式指令和输出预算，中途还修复了编码与内存问题，这些都保留了记录。')

s=base(7,'500-Question Results','Explicit Reasoning Helps the Action Loop','All four methods completed the same 500 questions • Scores recomputed from saved episodes')
names={'standard':'Standard','cot':'CoT','act':'Act-only','react':'ReAct'}
paper={'standard':28.7,'cot':29.4,'act':25.7,'react':27.4}
rows=[]
for m,rs in basic.items():
    rows.append([names[m],f'{paper[m]:.1f}%',f'{100*statistics.mean(r["em"] for r in rs):.1f}%',f'{100*statistics.mean(r["f1"] for r in rs):.2f}%'])
table(s,1.11,2.3,7.36,['Method','Paper EM','Our EM','Our F1'],rows,[2.04,1.77,1.77,1.78],.75,highlight=3,size=18)
panel(s,8.83,2.3,3.39,1.53,'E3F3ED')
text(s,9.03,2.5,2.99,.55,'+21.2 pp',30,GREEN,True)
text(s,9.03,3.13,2.99,.48,'ReAct vs. Act-only EM',15,GRAY)
panel(s,8.83,4.08,3.39,1.66)
text(s,9.03,4.27,2.99,.5,'Only 3 questions',21,BLUE,True)
text(s,9.03,4.91,2.99,.55,'Separate CoT and ReAct\n185 vs. 182 correct',15,GRAY)
text(s,1.11,6.28,11.11,.39,'Step-budget failures: Act-only 145 / 500 → ReAct 71 / 500',20,GREEN,True)
text(s,1.11,6.91,11.11,.32,'Different models and evaluation conditions: compare the pattern, not a claimed gain over the paper.',13,GRAY)
add_note(s,'第8页：500题结果',55,'四组都完成了相同的五百题，样本和顺序一致。ReAct 的 EM 为百分之三十六点四，比 Act-only 高二十一点二个百分点。Act-only 有一百四十五题在七步内未完成回答，ReAct 是七十一题，支持显式推理有助于规划检索和推进任务的解释，但不能把所有差距都归因于推理。CoT 答对一百八十五题，ReAct 答对一百八十二题，仅差三题，因此本次实验不能证明谁稳定更好。Standard 仍能利用模型的内部知识回答一部分问题。结果方向与论文一致：ReAct 优于 Act-only，但没有超过 CoT。模型和条件不同，绝对分数不能作为超过论文的证据。F1 是我们额外报告的词元重合指标。')

s=base(8,'50-Question Hybrid Exploration','A Higher Score Does Not Prove a Better Fallback','Seven methods on the same first 50 questions • Basic runs and hybrid runs were independent')
all50={**{m:rs[:50] for m,rs in basic.items()},**hybrid}
display={**names,'cot_sc':'CoT-SC','cot_sc_to_react':'CoT-SC → ReAct','react_to_cot_sc':'ReAct → CoT-SC'}
rows=[]
for m,rs in all50.items():
    rows.append([display[m],f'{sum(r["em"] for r in rs)}/50',f'{100*statistics.mean(r["em"] for r in rs):.0f}%',f'{sum(len(r["calls"]) for r in rs):,}'])
table(s,1.11,2.24,7.12,['Method','Correct','EM','Calls'],rows,[3.04,1.3,1.2,1.58],.51,highlight=6,size=15)
panel(s,8.61,2.24,3.61,1.63,'E3F3ED')
text(s,8.82,2.42,3.19,.34,'CoT-SC → ReAct',18,GREEN,True)
text(s,8.82,2.94,3.19,.73,'16 switches\n2 corrected; 1 made wrong\nNet benefit: +1 answer',15)
panel(s,8.61,4.08,3.61,1.94)
text(s,8.82,4.27,3.19,.34,'ReAct → CoT-SC',18,BLUE,True)
text(s,8.82,4.81,3.19,.93,'7 fallbacks; 0 rescued\n48% was already achieved\nby its fresh ReAct stage.',15)
text(s,1.11,6.62,11.11,.4,'Small sample: one answer = 2 percentage points. No stable hybrid advantage established.',17,GREEN,True)
text(s,1.11,7.1,11.11,.22,'CoT-SC uses 21 samples per question. Independent generation and live retrieval can vary across runs.',11,GRAY)
add_note(s,'第9页：50题组合探索',60,'我们还在同样的前五十题上完成了三种新增方法，与之前四组的前五十题比较。CoT-SC 没有超过普通 CoT，却使用了二十一倍调用次数。CoT-SC 转 ReAct 在十六题上切换，纠正两题，也把一题原本正确的答案改错，净增加一道正确答案。反向组合总成绩达到百分之四十八，看起来最高，但检查轨迹后发现，七次回退没有救回任何答案。这一组重新运行的 ReAct 阶段本身已经答对二十四题，因此不能把百分之四十八归因于回退策略。独立生成和实时检索仍会波动。五十题中一道题占两个百分点，我们目前确认了组合流程能够执行，但没有建立稳定的组合收益。')

s=base(9,'Conclusion','Mechanism Reproduced; Exact Replication Remains Open','Assessing what this reproduction demonstrates — and what it does not')
cards=[
    (1.11,'REPRODUCED','Prompt ablations and\nWikipedia feedback loop','4 methods × 500 questions\n7 methods compared on 50',GREEN,'E3F3ED'),
    (4.9,'OBSERVED','ReAct > Act-only\nReAct ≈ CoT','Hybrid benefit is limited\nand sensitive to run variation.',BLUE,'F0F4FA'),
    (8.69,'LIMITS','Substitute model\nAdapted implementation','Single seed; live Wikipedia\nOther tasks not reproduced.',GRAY,'F0F4FA'),
]
for x,label,headline,body,color,bg in cards:
    panel(s,x,2.36,3.53,3.47,bg)
    text(s,x+.22,2.61,3.09,.3,label,13,color,True)
    text(s,x+.22,3.18,3.09,1.02,headline,21,NAVY,True)
    text(s,x+.22,4.58,3.09,.95,body,16,GRAY)
text(s,1.11,6.24,11.11,.66,'Next: shared-stage hybrid comparisons, more seeds, and targeted failure analysis.',20,GREEN,True)
text(s,1.11,7.03,11.11,.24,'Evaluation: evidence for explicit reasoning in acting; no claim of full-paper or original-model replication.',12,GRAY)
add_note(s,'第10页：复现评估',35,'最后评估这次复现。我们完成了提示消融、模型与 Wikipedia 的反馈循环，以及四组五百题和七组五十题的比较。结果支持显式推理有助于行动，但没有证明 ReAct 稳定优于 CoT，也没有证明组合方法稳定获益。我们的模型、聊天接口、输出预算和环境修复都与原实验存在差异，而且目前只有一个随机种子，其他任务尚未复现。下一步应固定或共享组合方法的第一阶段输出，隔离回退策略的贡献，再用更多种子和失败分析验证结论。因此，这次是方法机制的复现，而不是原模型分数或整篇论文的完整复现。')

for s in list(prs.slides)[5:]:
    for x in s.shapes:
        assert x.left>=0 and x.top>=0
        assert x.left+x.width<=prs.slide_width and x.top+x.height<=prs.slide_height
prs.save(OUTPUT)
(ROOT/'ReAct_final_reproduction_notes.md').write_text('# 复现部分讲稿\n\n第6–10页，约4分钟。前5页建议约3分半钟，总时长约7分半钟。\n\n'+'\n'.join(notes),encoding='utf-8')
check=Presentation(OUTPUT)
assert len(check.slides)==10
print(f'Saved {OUTPUT.name}: 10 slides; five revised slides with speaker notes; verified sample alignment and slide bounds.')
