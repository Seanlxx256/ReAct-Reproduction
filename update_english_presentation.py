"""Preserve the user's latest deck, improve slide 2, and export English notes."""
from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parent
SOURCE=ROOT/'ReAct_HotpotQA_Reproduction_final.pptx'
OUT=ROOT/'ReAct_HotpotQA_Reproduction_English.pptx'
WORD=ROOT/'ReAct_English_Presentation_Script.docx'
prs=Presentation(SOURCE)
scripts=[
('Knowledge-Intensive Tasks', '''The paper studies two tasks that need facts and reasoning. HotpotQA asks questions that usually require information from more than one Wikipedia page. FEVER gives the model a claim. The model must decide whether evidence supports it, contradicts it, or is not enough to make a decision.
The goal is to see whether reasoning and outside information work better together. The authors measure answer accuracy and also examine factual mistakes. Today, I will focus on HotpotQA and our reproduction of its experiments.'''),
('The Wikipedia Environment', '''For each test question, the model gets no supporting passages and no correct answer. It can use what it already knows. Act-only and ReAct can also search Wikipedia.
The interface has three actions. Search returns the first five sentences of an entity page, or suggests five similar entities. Lookup returns the next sentence on the current page that contains a keyword. Finish submits the answer.
The model chooses what to search. The program runs the action and returns an observation. ReAct adds that observation to the history before the next call. Thoughts and actions come from the model; observations come from the environment.'''),
('Four Basic Methods', '''The four basic methods differ in two ways: whether they show reasoning and whether they can retrieve information.
Standard gives an answer directly. CoT writes out its reasoning and then answers. Both use one model call and have no external search.
Act-only chooses actions and receives observations in a loop. It does not write out thoughts, although reasoning may still happen inside the model. ReAct adds explicit thoughts before actions. These thoughts, actions, and observations stay in the history for later steps.'''),
('Building the Few-Shot Prompts', '''The authors choose six training examples and write ReAct traces by hand. These examples guide the model without changing its parameters.
Here is an original example: which magazine started earlier, Arthur's Magazine or First for Women? The model plans to find both dates. It searches Arthur's Magazine and gets 1844. It then searches First for Women and gets 1989. Since 1844 is earlier, it submits Arthur's Magazine. The observations shown here are shortened.
The authors remove parts of these same examples to build the other prompts. Standard keeps questions and answers. CoT keeps reasoning too. Act-only keeps actions and observations but removes thoughts. ReAct keeps the full trace. Using the same source cases reduces differences caused by example choice.'''),
('The Paper’s Results', '''The main table uses PaLM-540B, a language model with 540 billion parameters. HotpotQA is scored with exact match, or EM: an answer must match the reference after basic text normalization.
ReAct scores 27.4, above Act-only at 25.7, but below CoT at 29.4. Search does not always improve accuracy. CoT can invent facts; ReAct can make reasoning mistakes or fail to find useful information.
CoT-SC samples 21 answers at temperature 0.7 and takes the most common answer. The combined methods switch when ReAct runs out of steps, or when the CoT vote is below half the samples. The curve shows how scores change with more samples. ReAct to CoT-SC reaches the best prompting score, 35.1. The supervised result uses different training conditions.'''),
('Our Reproduction Design', '''We use the same question-only task and the original six examples. We replace the original model with a fixed version of GPT-4.1 mini. This tests the method with a different model; it is not an exact repeat of the original experiment.
The four basic methods use the same 500 questions, with random seed 233. The three added methods use the first 50 of those questions. Basic methods use temperature zero and up to seven interaction steps. CoT-SC uses 21 samples at temperature 0.7.
We do not enable built-in search or an extra reasoning mode. That does not mean the model has no internal reasoning.'''),
('How the Scripts Work', '''We turn the original notebook into scripts for batch runs. Run QA builds prompts, schedules questions, and saves results. Run HotpotQA calls the model and reads its actions. WikiEnv searches Wikipedia. Wrappers loads questions and calculates scores.
For Act-only and ReAct, each observation goes back into the history for the next call. Standard and CoT do not enter this loop.
We first checked small samples, then ran 100 questions. Extend QA expanded the target to 500 without repeating completed questions. We save raw outputs, full traces, scores, and token usage. We also adapted the chat format and output limits, and recorded fixes for text encoding and memory errors.'''),
('Our 500-Question Results', '''All four methods completed the same 500 questions. ReAct scores 36.4 percent EM, compared with 15.2 percent for Act-only: a gain of 21.2 percentage points. Act-only failed to finish within seven steps on 145 questions. ReAct did so on 71. This supports the value of explicit reasoning during action, though it does not explain every score difference.
CoT scores 37 percent. It answers 185 questions correctly, while ReAct answers 182. A three-question gap is too small to claim a clear winner here. F1 measures word overlap and gives partial credit.
The pattern agrees with the paper, but different models and conditions prevent a direct claim that we beat its results.'''),
('The 50-Question Hybrid Test', '''We compare all seven methods on the same first 50 questions. CoT-SC matches CoT's 42 percent EM but uses 21 times as many calls.
CoT-SC to ReAct switches on 16 questions. It fixes two answers and changes one correct answer to a wrong one: a net gain of one.
The reverse method scores 48 percent, but none of its seven fallbacks rescues an answer. Its new ReAct stage already gets 24 questions right. So the higher score is not evidence that the fallback helps. Separate runs can vary, even at temperature zero.
One answer changes the score by two percentage points. This sample does not show a stable hybrid advantage.'''),
('Conclusion', '''We reproduced the prompt comparisons and the Wikipedia feedback loop. We completed four methods on 500 questions and compared seven methods on 50.
Our results support explicit reasoning during action. They do not show that ReAct clearly beats CoT, or that hybrid methods reliably improve results.
The limits include a different model, one random seed, live Wikipedia, and changes to the implementation. We have not reproduced the paper's other tasks. A useful next step is to reuse the same first-stage outputs when testing fallbacks, then test more seeds. This is a reproduction of the method, not of the entire paper or its exact scores.'''),
]
assert len(prs.slides)==len(scripts)
for slide,(_,script) in zip(prs.slides,scripts):
    slide.notes_slide.notes_text_frame.text=script

s=prs.slides[1]
for x in list(s.shapes):x._element.getparent().remove(x._element)
s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('FFFFFF')
def text(x,y,w,h,value,size=18,color='0D1326',bold=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name='Noto Sans SC';p.font.size=Pt(size)
        p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(8)
def panel(x,y,w,h,color):
    box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    box.fill.solid();box.fill.fore_color.rgb=RGBColor.from_string(color);box.line.fill.background()
text(1.11,.44,11.11,.33,'02 · HotpotQA Environment',14,'235BC4',True)
text(1.11,.88,11.11,.65,'Three Actions. One Wikipedia Feedback Loop.',28,bold=True)
text(1.11,1.55,11.11,.4,'Question-only input • The model decides what to retrieve',17,'64748B')
cards=[(1.11,'01  SEARCH','Search[entity]','Page found\nFirst 5 sentences\n\nNot found\nTop 5 similar entities','235BC4'),
       (4.9,'02  LOOKUP','Lookup[string]','Find the next sentence\ncontaining the keyword\non the current page.\n\nLike Ctrl+F.','147D69'),
       (8.69,'03  FINISH','Finish[answer]','Submit the final answer\nand end the task.','235BC4')]
for x,label,action,body,color in cards:
    panel(x,2.18,3.53,2.94,'F0F4FA')
    text(x+.22,2.4,3.09,.26,label,13,color,True)
    text(x+.22,2.96,3.09,.43,action,21,bold=True)
    text(x+.22,3.56,3.09,1.42,body,16,'475569')
text(1.11,5.5,11.11,.26,'WHO PRODUCES WHAT?',13,'235BC4',True)
for x,title,body,color in [(1.11,'MODEL','Thought + Action','F0F4FA'),(4.9,'ENVIRONMENT','Observation','E3F3ED'),(8.69,'NEXT CALL','History + new observation','F0F4FA')]:
    panel(x,5.96,3.53,.8,color)
    text(x+.18,6.08,3.17,.2,title,11,'64748B',True)
    text(x+.18,6.39,3.17,.28,body,16,bold=True)
text(4.64,6.2,.25,.32,'→',18,'64748B')
text(8.43,6.2,.25,.32,'→',18,'64748B')
text(1.11,7.03,11.11,.25,'Restricted retrieval: no supporting passages are supplied with the test question.',12,'64748B')
for x in s.shapes:
    assert x.left+x.width<=prs.slide_width and x.top+x.height<=prs.slide_height
prs.save(OUT)

# A standard Office Open XML document, with no extra package dependency.
def paragraph(value,style=None):
    prop=f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ''
    return f'<w:p>{prop}<w:r><w:t xml:space="preserve">{escape(value)}</w:t></w:r></w:p>'
content=[paragraph('ReAct: HotpotQA Reproduction','Title'),paragraph('English presentation script | 10 slides | About 8 minutes','Subtitle')]
for i,(title,script) in enumerate(scripts,1):
    content.append(paragraph(f'Slide {i} — {title}','Heading1'))
    content.extend(paragraph(p) for p in script.split('\n'))
body=''.join(content)
document=f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{body}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1134"/></w:sectPr></w:body></w:document>'''
styles='''<?xml version="1.0" encoding="UTF-8"?><w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="23"/></w:rPr></w:rPrDefault><w:pPrDefault><w:pPr><w:spacing w:after="140" w:line="276" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults><w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style><w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:pPr><w:keepNext/></w:pPr><w:rPr><w:b/><w:color w:val="0D1326"/><w:sz w:val="40"/></w:rPr></w:style><w:style w:type="paragraph" w:styleId="Subtitle"><w:name w:val="Subtitle"/><w:basedOn w:val="Normal"/><w:rPr><w:color w:val="64748B"/><w:sz w:val="24"/></w:rPr></w:style><w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:pPr><w:keepNext/><w:spacing w:before="280" w:after="140"/><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:b/><w:color w:val="235BC4"/><w:sz w:val="28"/></w:rPr></w:style></w:styles>'''
types='''<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>'''
rels='''<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'''
docrels='''<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>'''
with ZipFile(WORD,'w',ZIP_DEFLATED) as z:
    for name,value in {'[Content_Types].xml':types,'_rels/.rels':rels,'word/document.xml':document,'word/styles.xml':styles,'word/_rels/document.xml.rels':docrels}.items():z.writestr(name,value)
check=Presentation(OUT)
for slide,(_,script) in zip(check.slides,scripts):
    assert slide.notes_slide.notes_text_frame.text==script
    assert not re.search('[\u4e00-\u9fff]',script)
from xml.etree import ElementTree as ET
with ZipFile(WORD) as z:
    assert z.testzip() is None
    for name in z.namelist():ET.fromstring(z.read(name))
    extracted=' '.join(ET.fromstring(z.read('word/document.xml')).itertext())
    assert all(script.split('\n')[0] in extracted for _,script in scripts)
print('Saved:',OUT.name,WORD.name)
print('Verified: 10 matching English slide notes; valid DOCX XML and ZIP; slide 2 within bounds.')
print('Script words:',sum(len(script.split()) for _,script in scripts))
