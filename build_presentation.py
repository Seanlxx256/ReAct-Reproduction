"""Build a Chinese 7-minute ReAct presentation from the saved experiment records."""
from datetime import datetime
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
BG = '101B2D'
PANEL = '1C2B43'
WHITE = 'F3F6FC'
MUTED = 'A8B8CF'
CYAN = '49D9CF'
GOLD = 'F6C76A'
FONT = 'Microsoft YaHei'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
notes = []


def rgb(code):
    return RGBColor.from_string(code)


def box(slide, x, y, w, h, color=PANEL, rounded=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    return s


def text(slide, x, y, w, h, value, size=22, color=WHITE, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.03)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(10)
    return shape


def slide(title, subtitle, source, speaker, seconds):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    box(s, .6, .45, .09, .7, CYAN, False)
    text(s, .85, .45, 11.8, .7, title, 30, bold=True)
    text(s, .85, 1.2, 11.8, .55, subtitle, 17, MUTED)
    text(s, .65, 7.04, 11.8, .25, source, 10, MUTED)
    text(s, 12.35, 7.02, .5, .3, f'{len(prs.slides):02d}', 12, CYAN)
    s.notes_slide.notes_text_frame.text = f'建议时长：{seconds} 秒\n\n{speaker}'
    notes.append((title, seconds, speaker, source))
    return s


def table(s, headers, rows, x=.8, y=2, w=11.7, row_h=.47, widths=None):
    shape = s.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y), Inches(w),
                               Inches(row_h*(len(rows)+1)))
    t = shape.table
    if widths:
        for col, frac in zip(t.columns, widths):
            col.width = Inches(w*frac)
    for r, values in enumerate([headers]+rows):
        for c, value in enumerate(values):
            cell = t.cell(r, c)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb('263C57' if r == 0 else PANEL)
            cell.margin_left = Inches(.12)
            cell.margin_top = Inches(.06)
            cell.margin_bottom = Inches(.03)
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(15 if w < 4 else (17 if len(rows)>6 else 19))
                p.font.bold = r == 0
                p.font.color.rgb = rgb(CYAN if r == 0 else WHITE)
    return shape


methods = ['standard', 'cot', 'act', 'react']
labels = ['Standard', 'CoT', 'Act-only', 'ReAct']
hundred = {}
progress = {}
indices = None
for m in methods:
    folder = ROOT / 'results/qa-v5-100/hotpotqa' / m
    records = [json.loads(line) for line in (folder/'episodes.jsonl').read_text(encoding='utf-8').splitlines()]
    assert len(records) >= 100
    subset = records[:100]
    ids = [r['question_idx'] for r in subset]
    assert indices is None or indices == ids
    indices = ids
    hundred[m] = {'em': sum(r['em'] for r in subset)/100,
                  'f1': sum(r['f1'] for r in subset)/100,
                  'calls': sum(len(r['calls']) for r in subset),
                  'forced': sum(st.get('forced_finish', False) for r in subset for st in r['stages']),
                  'repairs': sum(st.get('n_badcalls', 0) for r in subset for st in r['stages']),
                  'truncated': sum(c['finish_reason']=='length' for r in subset for c in r['calls'])}
    progress[m] = len(records)
twenty = json.loads((ROOT/'results/qa-v4-20/hotpotqa/comparison.json').read_text(encoding='utf-8'))

s = slide('ReAct：推理与行动如何形成闭环？', '论文实验解读与 HotpotQA 复现  ·  6–8 分钟 presentation',
          'Yao et al., ICLR 2023 · ReAct.pdf · 本地实验记录快照：2026-10-07',
          '今天介绍 ReAct 如何把显式推理和外部行动结合起来。我先说明原论文怎样设计对照，以及它究竟得到了哪些结果；再介绍我们如何用现有模型复现这一机制，最后分析实际成绩和局限。这里的复现采用替代模型，不能理解为复现原模型的精确分数。', 25)
text(s, .85, 2.1, 11.5, 1.4, 'Thought → Action → Observation', 42, CYAN, True)
text(s, .9, 3.7, 11.3, 1.3, '核心问题：外部反馈是否改善回答？\n同样允许检索时，显式推理是否改善行动？', 26)
text(s, .9, 5.6, 11.3, .7, '已完成：七组 × 20 条；四组共同完成前 100 条', 23, GOLD)

s = slide('01  方法：让模型在行动后重新思考', '模型决定查什么；程序执行请求；真实观察进入下一轮输入',
          '论文 §2–3；实现：run_hotpotqa.py / wikienv.py',
          'ReAct 并不是一次性输出完整推理后结束。模型生成本轮 Thought 和 Action，脚本执行 Search 或 Lookup，并把真实结果写成 Observation。下一轮模型同时看到过去的推理、动作和观察，再调整计划。示例中的观察是固定教学材料，当前问题的观察则来自实际 Wikipedia 请求。模型不负责伪造观察。', 40)
for x, heading, body in [( .85, 'Thought', '拆解问题\n解释反馈\n调整下一步计划'), (4.95, 'Action', 'Search[实体]\nLookup[关键词]\nFinish[答案]'), (9.05, 'Observation', '真实 Wikipedia 内容\n追加到上下文\n供下一轮使用')]:
    box(s, x, 2.2, 3.45, 2.65)
    text(s, x+.18, 2.42, 3.1, .5, heading, 28, CYAN, True)
    text(s, x+.18, 3.15, 3.05, 1.5, body, 21)
text(s, 4.38, 3.2, .5, .7, '→', 30, GOLD)
text(s, 8.48, 3.2, .5, .7, '→', 30, GOLD)
text(s, 1, 5.35, 11.3, .9, '下一轮输入 = 示例 + 问题 + 过去的 Thought / Action / Observation', 23)

s = slide('02  原论文：四类任务与受控对照', 'QA 检验“行动辅助推理”；交互任务检验“推理指导行动”',
          'ReAct.pdf §3–4，pp.4–8；500 条 GPT-3 设置见附录 A.1，p.14',
          '论文覆盖四类任务。HotpotQA 是多跳问答，FEVER 是事实核查；模型只看到问题或声明，不拿到支持段落。示例分别为六条和三条。ALFWorld 是六类家庭任务，在 134 个未见游戏上评估；WebShop 是 500 条购物指令。主模型是 PaLM-540B，附录还有 GPT-3 实验。我们沿用的 500 条来自附录和仓库；200 条则是人工轨迹分析，不是准确率评估规模。', 45)
table(s, ['任务', '输入与环境', '评价指标'], [
    ['HotpotQA', '问题 → Wikipedia；6-shot', '答案 EM'],
    ['FEVER', '声明 → Wikipedia；3-shot', '标签准确率'],
    ['ALFWorld', '官方文字引擎；134 个未见任务', '任务成功率'],
    ['WebShop', '购物网站；500 条测试指令', '平均得分 / 成功率']], row_h=.6)
text(s, 1, 5.5, 11.2, .9, 'QA 对照：Standard / CoT / Act / ReAct\n扩展：CoT-SC（21 次投票）与两种切换策略', 22, CYAN)

s = slide('03  论文结果：优势取决于任务与组合', '主表采用 PaLM-540B；原始 ReAct 并非在所有 QA 上最好',
          'ReAct.pdf Table 1（p.5）、Table 3 / 4（p.8）；均为百分数',
          '主表中 ReAct 的 HotpotQA EM 是 27.4，低于 Standard 和 CoT，但在 FEVER 上优于它们。组合方法表现更好：ReAct 转 CoT-SC 在 HotpotQA 达到 35.1，CoT-SC 转 ReAct 在 FEVER 达到 64.6。交互任务中，ALFWorld 的 ReAct 最佳提示组合约 71%，高于 Act 的 45%；WebShop 成功率为 40%，高于 Act 的 30.1%。这些结果支持方法互补，而不是联网或显式推理必然提高分数。', 55)
table(s, ['方法', 'HotpotQA EM', 'FEVER Acc'], [
    ['Standard', '28.7', '57.1'], ['CoT', '29.4', '56.3'], ['CoT-SC', '33.4', '60.4'],
    ['Act', '25.7', '58.9'], ['ReAct', '27.4', '60.9'],
    ['CoT-SC → ReAct', '34.2', '64.6'], ['ReAct → CoT-SC', '35.1', '62.0']],
    x=.85, y=1.95, w=8.1, row_h=.49, widths=[.52,.24,.24])
box(s, 9.3, 2, 3.15, 3.85)
text(s, 9.48, 2.25, 2.8, .5, '交互任务', 25, CYAN, True)
text(s, 9.48, 3, 2.8, 2.7, 'ALFWorld（best of 6）\nAct 45 → ReAct 71\n\nWebShop SR\nAct 30.1 → ReAct 40.0', 19)
text(s, .95, 6.15, 11.5, .45, '外部信息更有依据；内部知识与检索的结合更有潜力。', 22, GOLD)

s = slide('04  我们怎样复现：统一条件，改变提示方法', '固定模型与样本；标准答案仅用于评分；不提供支持段落',
          '本地配置：results/qa-v5-100/hotpotqa/*/config.json；提示：prompts_naive.json',
          '我们使用固定版本 GPT-4.1 mini 替代已经停用的 text-davinci-002。保留仓库的六个示例、数据、检索动作和评分规则，以种子 233 固定抽样，各方法使用同样的问题。全部采用温度零。Act 和 ReAct 最多七轮，每次最多 256 输出 token；CoT 最多 512。GPT-4.1 mini 没有本次启用的专门 thinking 通道，但这不代表内部完全没有推理。比较的变量是显式轨迹和外部交互。', 45)
table(s, ['统一项', '实际设置'], [
    ['模型', 'gpt-4.1-mini-2025-04-14'],
    ['数据 / 抽样', 'HotpotQA dev；seed=233；相同索引'],
    ['提示与解码', '6-shot；temperature=0；question-only'],
    ['交互预算', '最多 7 轮；Act/ReAct 输出上限 256 token'],
    ['CoT 与评分', 'CoT 上限 512 token；原始 EM + 辅助 F1']], row_h=.55, widths=[.25,.75])
text(s, .95, 5.8, 11.3, .7, '实现偏差：替换模型、聊天格式指令、输出上限；结果不能直接对齐原 PaLM 数值。', 21, GOLD)

s = slide('05  具体实现：Prompt 与执行路径', '同一 API 模型，四种输入轨迹；网络访问只由脚本执行',
          'run_qa.py → run_hotpotqa.py / wrappers.py / wikienv.py',
          'Standard 的示例只有问题和答案，末尾为 Answer。CoT 增加 Thought，末尾要求开始逐步推理。Act 的示例有动作和观察，末尾为 Action；ReAct 额外包含 Thought。当前问题的观察来自 WikiEnv 请求 Wikipedia，并进入后续输入。我们还保存原始输出、token 用量和文件哈希，支持断点续跑。修复了请求 403 与解析缺陷，并把原来的 100 token 上限提高到 256，这个变化已在配置中记录。', 45)
table(s, ['方法', '示例结构 / 请求末尾', '外部检索'], [
    ['Standard', 'Question → Answer；末尾 Answer:', '无'],
    ['CoT', 'Question → Thought → Answer；末尾 Thought:', '无'],
    ['Act-only', 'Action → Observation；末尾 Action 1:', '有'],
    ['ReAct', 'Thought → Action → Observation；末尾 Thought 1:', '有']], row_h=.61, widths=[.19,.64,.17])
text(s, .95, 5.55, 11.3, 1, '模型原始输出 → 动作解析 → Wikipedia → 观察回填\n每题保存 episodes.jsonl；逐题评分，summary.json 取平均。', 22, CYAN)

s = slide('06  公平对照：共同完成的前 100 条', '四组同题比较；图示 EM，表格补充 F1 与交互失败',
          '数据：qa-v5-100/hotpotqa/*/episodes.jsonl 的前 100 条；非未完成的 500 条混合汇总',
          '这页只比较四组共同完成的前一百条。Standard 为 37%，CoT 为 39%，Act 为 18%，ReAct 为 37%。ReAct 比 Act 高十九个百分点，而且步数耗尽从二十七条降到十九条；但没有超过 Standard，略低于 CoT。CoT 和 ReAct 仅差两道题，不足以证明稳定差异。没有截断或无效动作，ReAct 三次格式补救成功。Standard 高分说明已有知识能回答部分问题；检索也可能带来实体选择和理解错误。', 60)
for i, (m, label) in enumerate(zip(methods, labels)):
    y=2.15+i*.86
    text(s, .9, y, 2, .45, label, 22)
    box(s, 3, y+.05, 5.4, .34, PANEL, False)
    box(s, 3, y+.05, 5.4*hundred[m]['em']/.5, .34, CYAN if m=='react' else '6588B2', False)
    text(s, 8.6, y-.02, 1, .5, f"{hundred[m]['em']:.0%}", 25, CYAN if m=='react' else WHITE, True)
table(s, ['方法', 'F1', '步数耗尽'], [[label, f"{hundred[m]['f1']*100:.2f}%", str(hundred[m]['forced']) if m in ('act','react') else '—'] for m,label in zip(methods,labels)],
      x=9.75, y=2.05, w=2.85, row_h=.62, widths=[.42,.28,.30])
text(s, .95, 6.1, 11.4, .6, 'ReAct − Act：+19 pp；ReAct − CoT：−2 pp。结果描述，不代表显著性结论。', 21, GOLD)

s = slide('07  组合方法：20 条探索性结果', '投票一致不等于正确；切换检索可能补救，也可能丢失正确答案',
          '数据：results/qa-v4-20/hotpotqa/comparison.json 与各组轨迹；各方法独立采样',
          '我们另完成了七组各二十条的冒烟测试。CoT-SC 每题采样二十一条，组合方法按原论文规则切换。CoT-SC 转 ReAct 得到 55%，只比 50% 多一道题，不能宣称稳定提升。六题切换中，检索修正了两题，但也丢失了两题原本正确的投票答案。反向组合三题触发回退，没有救回正确答案。高度一致的错误投票也存在。这说明切换策略有意义，但不保证成功。', 35)
table(s, ['方法', 'EM', 'API 调用'], [[label,f"{twenty[m]['em']*100:.0f}%",str(twenty[m]['n_calls'])] for m,label in [
    ('cot','CoT'),('react','ReAct'),('cot_sc','CoT-SC'),('cot_sc_to_react','CoT-SC → ReAct'),('react_to_cot_sc','ReAct → CoT-SC')]],
    x=.85,y=2,w=7.3,row_h=.58,widths=[.6,.2,.2])
box(s,8.6,2,3.85,3.8)
text(s,8.8,2.25,3.4,.6,'切换观察',26,CYAN,True)
text(s,8.8,3.05,3.4,2.5,'CoT-SC → ReAct\n6 题触发：2 题修正\n2 题失去原正确答案\n\nReAct → CoT-SC\n3 题触发：均未救回',20)

s = slide('08  错误分析：能力、格式与指标要分开', '没有把“看起来合理”自动改判为正确，也没有针对测试题提供提示',
          '轨迹示例：qa-v4-20 / qa-v5-100；论文 Table 2 同样讨论 reasoning/search/label ambiguity',
          '我们区分三种问题。第一是模型能力：例如把电影 November 当成月份，导致搜索演员生日；不能专门告诉模型答案来修分。第二是接口和格式：初版截断动作并把观察当成动作，这应修复，新版已没有无效动作。第三是评分局限：回答 Psych 而标准标签为一整句，EM 仍是零；Act 输出长句也容易失分。F1 提供部分重合信息，但不是真正的语义裁判。', 45)
for y,title,body in [
    (2.0,'能力失败','把《November》当月份；检索实体错误；超过步数预算。保留失败。'),
    (3.45,'工程缺陷','初版 100-token 截断、误解析 Observation。修复，并记录实现变化。'),
    (4.9,'评估局限','Psych 与长标准答案不完全匹配；Act 的冗长回答也会导致 EM 失分。')]:
    box(s,.85,y,11.6,1.2)
    text(s,1.05,y+.12,2.05,.6,title,24,CYAN,True)
    text(s,3.15,y+.18,9, .8,body,21)

s = slide('09  结论与下一步：方法复现，尚非完整复现', '当前证据支持运行机制与局部对照；仍需扩大共同样本与跨任务验证',
          '实验快照：2026-10-07；原论文：Yao et al., ReAct, ICLR 2023',
          '目前完成的是替换模型后的方法复现。共同一百条上，ReAct 优于 Act，但没有超过 CoT 或 Standard。二十条的组合结果提示互补可能，却不足以作强结论。五百条扩展尚未全部结束，不能把不同样本数的成绩混在一起比较。下一步先完成四组相同五百条，再评估 FEVER，并做更系统的错误分析。我们尚未做官方 ALFWorld、WebShop、微调和完整人工标注。',35)
text(s,.95,2,11.4,1.25,'已验证：真实检索闭环、显式轨迹回填、统一评分与可恢复记录。\n观察到：ReAct 优于 Act；与内部知识基线相比优势不稳定。',25)
table(s,['500 条扩展进度','Standard','CoT','Act-only','ReAct'],
      [['已完成样本',*[str(progress[m])+'/500' for m in methods]]],y=3.8,row_h=.65,widths=[.28,.18,.18,.18,.18])
text(s,.95,5.55,11.5,.85,'下一步：完成同题 500 条 → FEVER 对照 → 分层错误分析\n局限：替代模型、实时 Wikipedia、有限样本、聊天提示适配。',22,GOLD)

prs.core_properties.title = 'ReAct：论文实验与 HotpotQA 复现'
prs.core_properties.subject = '6–8 minute research presentation'
prs.core_properties.author = 'ReAct reproduction project'
prs.save(ROOT/'ReAct_presentation.pptx')
md = ['# ReAct presentation 讲稿', '', f'共 {len(notes)} 页；建议总时长 {sum(n[1] for n in notes)} 秒（约 7 分 10 秒）。',
      '', '数据快照：2026-10-07。100 条图表由前 100 条轨迹重算，避免混用尚未完成的 500 条汇总。', '']
for i,(title,seconds,speaker,source) in enumerate(notes,1):
    md += [f'## 第 {i} 页：{title}', '', f'建议 {seconds} 秒。', '', speaker, '', f'依据：{source}', '']
(ROOT/'ReAct_speaker_notes.md').write_text('\n'.join(md),encoding='utf-8')
(ROOT/'ReAct_presentation_data.json').write_text(json.dumps({'snapshot':'2026-10-07','hundred':hundred,
    'progress_500':progress,'twenty':twenty,'indices_100':indices},indent=2,ensure_ascii=False),encoding='utf-8')
print('Generated PPTX, speaker notes, data snapshot. Slides:',len(prs.slides),'Seconds:',sum(n[1] for n in notes))
